from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO


def export_rows_to_excel(rows: List[Dict[str, Any]], title: str = "Reporte") -> bytes:
    """
    Exporta una lista de diccionarios a un archivo Excel
    Retorna bytes del archivo Excel generado
    """
    # Defensiva: manejar casos edge
    if rows is None:
        rows = []
    # Si rows es un dict (error típico), lánzalo o conviértelo a lista segura:
    if isinstance(rows, dict) and "items" in rows:
        rows = rows["items"]
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    
    if not rows:
        # Si no hay datos, mostrar mensaje
        ws['A1'] = "Sin datos"
        ws['A1'].font = Font(bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')
    else:
        # Obtener encabezados del primer diccionario
        headers = list(rows[0].keys())
        
        # Escribir encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        # Escribir datos
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header)
                # Convertir fechas y otros tipos a string
                if value is not None:
                    if hasattr(value, 'isoformat'):  # datetime, date
                        value = value.isoformat()
                    else:
                        value = str(value)
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Auto-ajustar ancho de columnas
        for col in range(1, len(headers) + 1):
            column_letter = get_column_letter(col)
            max_length = 0
            
            # Calcular longitud máxima del contenido
            for row in ws[column_letter]:
                try:
                    if row.value:
                        max_length = max(max_length, len(str(row.value)))
                except:
                    pass
            
            # Establecer ancho mínimo y máximo
            adjusted_width = min(max(max_length + 2, 10), 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Guardar en BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output.getvalue()
